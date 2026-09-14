"""Document loaders: PDF, HTML, Markdown, TXT, DOCX.

Optional parsers are imported lazily so the package works (and the test suite
runs) without every format's dependency installed. A missing parser raises a
clear, actionable error rather than an ImportError at module load.
"""
from __future__ import annotations

import hashlib
import html
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

SUPPORTED_SUFFIXES: frozenset[str] = frozenset(
    {".pdf", ".html", ".htm", ".md", ".markdown", ".txt", ".text", ".docx"}
)

_SCRIPT_STYLE = re.compile(
    r"<(script|style|noscript|nav|footer|header|aside)\b[^>]*>.*?</\1>",
    re.IGNORECASE | re.DOTALL,
)
_HTML_TITLE = re.compile(r"<title[^>]*>(.*?)</title>", re.IGNORECASE | re.DOTALL)
_HTML_H1 = re.compile(r"<h1[^>]*>(.*?)</h1>", re.IGNORECASE | re.DOTALL)
_META = re.compile(
    r"""<meta\s+[^>]*?(?:name|property)\s*=\s*["']([^"']+)["'][^>]*?content\s*=\s*["']([^"']*)["']""",
    re.IGNORECASE,
)
_BLOCK_TAGS = re.compile(r"</?(p|div|br|li|tr|h[1-6]|section|article)\b[^>]*>", re.IGNORECASE)
_TAG = re.compile(r"<[^>]+>")
_WS_RUNS = re.compile(r"[ \t\x0b\f\r]+")
_BLANK_RUNS = re.compile(r"\n{3,}")


@dataclass(frozen=True)
class LoadedDocument:
    """Raw parsed output, before chunking."""

    text: str
    title: str
    source_path: str
    content_sha256: str
    author: str | None = None
    publication_date: datetime | None = None
    page_count: int | None = None

    @property
    def is_empty(self) -> bool:
        return not self.text.strip()


class UnsupportedFormatError(ValueError):
    pass


class ParserUnavailableError(RuntimeError):
    """A format's optional dependency is not installed."""


def _sha256(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8", errors="replace")).hexdigest()


def clean_text(raw: str) -> str:
    """Normalise whitespace without destroying paragraph structure."""
    text = raw.replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    text = _WS_RUNS.sub(" ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    return _BLANK_RUNS.sub("\n\n", text).strip()


def _parse_date(value: str | None) -> datetime | None:
    if not value:
        return None
    candidate = value.strip().replace("Z", "+00:00")
    for fmt in (None, "%Y-%m-%d", "%Y/%m/%d", "%d %B %Y", "%B %d, %Y", "%Y"):
        try:
            dt = datetime.fromisoformat(candidate) if fmt is None else datetime.strptime(candidate, fmt)
        except (ValueError, TypeError):
            continue
        return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)
    return None


# -- format handlers -------------------------------------------------------

def load_html(raw: str, source: str = "") -> LoadedDocument:
    body = _SCRIPT_STYLE.sub(" ", raw)

    title_match = _HTML_TITLE.search(raw) or _HTML_H1.search(raw)
    title = html.unescape(_TAG.sub("", title_match.group(1))).strip() if title_match else ""

    meta = {k.lower(): v for k, v in _META.findall(raw)}
    author = meta.get("author") or meta.get("article:author") or meta.get("citation_author")
    published = _parse_date(
        meta.get("article:published_time")
        or meta.get("citation_publication_date")
        or meta.get("date")
        or meta.get("dc.date")
    )

    body = _BLOCK_TAGS.sub("\n", body)
    body = _TAG.sub(" ", body)
    text = clean_text(html.unescape(body))
    return LoadedDocument(
        text=text,
        title=title or (Path(source).stem if source else "Untitled"),
        source_path=source,
        content_sha256=_sha256(text),
        author=author.strip() if author else None,
        publication_date=published,
    )


def load_markdown(raw: str, source: str = "") -> LoadedDocument:
    text = clean_text(raw)
    title = ""
    for line in text.split("\n"):
        if line.startswith("# "):
            title = line[2:].strip()
            break
    return LoadedDocument(
        text=text,
        title=title or (Path(source).stem if source else "Untitled"),
        source_path=source,
        content_sha256=_sha256(text),
    )


def load_plaintext(raw: str, source: str = "") -> LoadedDocument:
    text = clean_text(raw)
    first = next((ln.strip() for ln in text.split("\n") if ln.strip()), "")
    return LoadedDocument(
        text=text,
        title=(first[:120] if len(first) <= 120 else Path(source).stem) or "Untitled",
        source_path=source,
        content_sha256=_sha256(text),
    )


def load_pdf(path: Path) -> LoadedDocument:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise ParserUnavailableError(
            "PDF ingestion requires 'pypdf'. Install it with: pip install pypdf"
        ) from exc

    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "") for page in reader.pages]
    text = clean_text("\n\n".join(pages))
    info = reader.metadata or {}
    return LoadedDocument(
        text=text,
        title=(getattr(info, "title", None) or "").strip() or path.stem,
        source_path=str(path),
        content_sha256=_sha256(text),
        author=(getattr(info, "author", None) or "").strip() or None,
        publication_date=_parse_date(str(getattr(info, "creation_date", "") or "")),
        page_count=len(pages),
    )


def load_docx(path: Path) -> LoadedDocument:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise ParserUnavailableError(
            "DOCX ingestion requires 'python-docx'. Install it with: pip install python-docx"
        ) from exc

    document = docx.Document(str(path))
    parts: list[str] = []
    for para in document.paragraphs:
        if not para.text.strip():
            continue
        style = (para.style.name or "").lower() if para.style else ""
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            parts.append(f"{'#' * min(int(level), 6)} {para.text.strip()}")
        else:
            parts.append(para.text.strip())

    text = clean_text("\n\n".join(parts))
    core = document.core_properties
    return LoadedDocument(
        text=text,
        title=(core.title or "").strip() or path.stem,
        source_path=str(path),
        content_sha256=_sha256(text),
        author=(core.author or "").strip() or None,
        publication_date=core.created.replace(tzinfo=timezone.utc) if core.created else None,
    )


def load_text(raw: str, *, suffix: str = ".txt", source: str = "") -> LoadedDocument:
    """Load already-in-memory content. Used for web fetches and tests."""
    suffix = suffix.lower()
    if suffix in {".html", ".htm"}:
        return load_html(raw, source)
    if suffix in {".md", ".markdown"}:
        return load_markdown(raw, source)
    if suffix in {".txt", ".text"}:
        return load_plaintext(raw, source)
    raise UnsupportedFormatError(f"cannot load in-memory content with suffix {suffix!r}")


def load_path(path: str | Path) -> LoadedDocument:
    """Load a document from disk, dispatching on file extension."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"no such document: {path}")
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise UnsupportedFormatError(
            f"unsupported format {suffix!r}; supported: {sorted(SUPPORTED_SUFFIXES)}"
        )
    if suffix == ".pdf":
        return load_pdf(path)
    if suffix == ".docx":
        return load_docx(path)
    return load_text(
        path.read_text(encoding="utf-8", errors="replace"), suffix=suffix, source=str(path)
    )
